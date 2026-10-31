"""Đổi file Excel dữ liệu của SLM thành JSONL để huấn luyện và đo. Excel là nguồn, JSONL sinh ra từ Excel, không sửa tay.

Mỗi loại một sheet (mac_dinh, ho_so, kich_ban, kien_thuc, trung_tam, ngoai_pham_vi). Tên sheet chính là nhãn.
Mỗi dòng: đầu vào của SLM (câu hỏi, hồ sơ, câu hệ thống đang hỏi) và đầu ra mong muốn
(loại, nhãn phụ, thực thể; skill md với hồ sơ và kịch bản khi có hồ sơ). Dòng sai chuẩn thì báo sheet và số dòng.

Chạy trong thư mục advisor:
  python -m scripts.build_slm_jsonl
  python -m scripts.build_slm_jsonl --excel data/slm/du_lieu_slm.xlsx --out data/slm/du_lieu_slm.jsonl
"""

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas.slm import INTENTS as TYPES
from app.schemas.slm import PROFILE_KEYS, check_entities, check_labels
from app.services.chat_flow import QUESTIONS

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "slm"
DEFAULT_EXCEL = DATA_DIR / "du_lieu_slm.xlsx"
DEFAULT_OUT = DATA_DIR / "du_lieu_slm.jsonl"

# Loại có skill riêng của học sinh (chỉ khi đã có hồ sơ). Các loại khác dùng skill chung cố định
SKILL_TYPES = ("ho_so", "kich_ban")
# Đầu vào (giống nhau ở mọi loại) rồi đầu ra. Loại lấy từ tên sheet, id sinh theo sheet và thứ tự dòng
COLUMNS = ("cau_hoi", "ho_so", "cau_he_thong_dang_hoi", "thuc_the", "nhan_phu", "tap")
SKILL_COLUMN = "skill_md"  # chỉ có ở sheet của SKILL_TYPES
SPLITS = ("train", "val", "test")


def columns_of(sheet: str) -> tuple[str, ...]:
    return COLUMNS[:-1] + (SKILL_COLUMN, "tap") if sheet in SKILL_TYPES else COLUMNS


def as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def as_json(value: Any, where: str, column: str) -> dict[str, Any]:
    text = as_text(value)
    if not text:
        return {}
    try:
        data = json.loads(text)
    except ValueError as ex:
        raise ValueError(f"{where}: cột {column} không phải JSON hợp lệ") from ex
    if not isinstance(data, dict):
        raise ValueError(f"{where}: cột {column} phải là một object JSON")
    return data


# Một dòng trong sheet thành một bản ghi. Sai chuẩn thì ném ValueError có tên sheet và số dòng
def to_record(intent: str, raw: dict[str, Any], where: str, record_id: str) -> dict[str, Any]:
    text = {key: as_text(raw.get(key)) for key in (*COLUMNS, SKILL_COLUMN)}
    sub = text["nhan_phu"]
    if not text["cau_hoi"]:
        raise ValueError(f"{where}: thiếu cau_hoi")
    if text["tap"] not in SPLITS:
        raise ValueError(f"{where}: tap phải là một trong {SPLITS}, đang là '{text['tap']}'")
    pending = text["cau_he_thong_dang_hoi"]
    if pending and pending not in QUESTIONS.values():
        raise ValueError(f"{where}: cau_he_thong_dang_hoi không khớp câu nào trong chat_flow.QUESTIONS")
    profile = as_json(raw.get("ho_so"), where, "ho_so")
    if set(profile) - set(PROFILE_KEYS):
        raise ValueError(f"{where}: ho_so có khóa lạ: {sorted(set(profile) - set(PROFILE_KEYS))}")
    entities = as_json(raw.get("thuc_the"), where, "thuc_the")
    # Cùng luật với app/schemas/slm.py, nên dữ liệu luôn khớp với đầu ra thật của SLM
    try:
        check_labels(intent, sub or None)
        check_entities(intent, sub or None, entities)
    except ValueError as ex:
        raise ValueError(f"{where}: {ex}") from ex
    # mac_dinh rỗng chỉ hợp lệ khi đang có câu chờ (tin mơ hồ, hỏi lại). Không có câu chờ thì phải có giá trị hồ sơ,
    # không thì là ngoai_pham_vi/mo_ho
    if intent == "mac_dinh" and not pending and not entities:
        raise ValueError(f"{where}: mac_dinh không có câu chờ thì thuc_the phải có giá trị hồ sơ")
    # Thành phố ghi đúng như học sinh viết, để SLM học tách nguyên văn chứ không tự đổi tên
    if "thanh_pho" in entities and entities["thanh_pho"] not in text["cau_hoi"]:
        raise ValueError(f"{where}: thanh_pho '{entities['thanh_pho']}' phải là đoạn chữ có trong cau_hoi")
    skill = text[SKILL_COLUMN]
    if skill and intent not in SKILL_TYPES:
        raise ValueError(f"{where}: loại {intent} không có skill riêng")
    if skill and not profile:
        raise ValueError(f"{where}: có skill_md thì phải có hồ sơ (skill sinh theo hồ sơ)")
    return {
        "id": record_id,
        "type": intent,
        "split": text["tap"],
        "input": {"cau_hoi": text["cau_hoi"], "ho_so": profile or None, "cau_he_thong_dang_hoi": pending},
        "output": {"type": intent, "nhan_phu": sub or None, "thuc_the": entities},
        "skill_md": skill or None,
    }


def read_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        if sheet.title not in TYPES:
            raise ValueError(f"Sheet '{sheet.title}' không phải một loại trong {TYPES} (mỗi loại một sheet)")
        rows = sheet.iter_rows(values_only=True)
        header = tuple(as_text(cell) for cell in next(rows, ()) if as_text(cell))
        if header != columns_of(sheet.title):
            raise ValueError(
                f"Sheet '{sheet.title}' phải có đúng các cột {columns_of(sheet.title)}, đang là {header}"
            )
        count = 0
        for number, values in enumerate(rows, start=2):
            raw = dict(zip(header, values, strict=False))
            if not any(as_text(v) for v in raw.values()):
                continue  # dòng trống
            count += 1
            where = f"{sheet.title}, dòng {number}"
            records.append(to_record(sheet.title, raw, where, f"{sheet.title}_{count:03d}"))
    return records


def write_jsonl(records: list[dict[str, Any]], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    records = read_records(args.excel)
    write_jsonl(records, args.out)
    print(f"Đã ghi {len(records)} dòng vào {args.out}")
    for intent in TYPES:
        rows = [r for r in records if r["type"] == intent]
        by_split = Counter(r["split"] for r in rows)
        skills = sum(1 for r in rows if r["skill_md"])
        print(
            f"  {intent:14} {len(rows):4}  "
            + "  ".join(f"{s}={by_split[s]}" for s in SPLITS)
            + (f"  có skill={skills}" if intent in SKILL_TYPES else "")
        )


if __name__ == "__main__":
    main()
