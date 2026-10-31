"""Dữ liệu của SLM: đủ 6 loại, đúng chuẩn, không tự mâu thuẫn, và JSONL khớp với Excel."""

import json
from collections import Counter, defaultdict

import pytest
from openpyxl import Workbook

from scripts.build_slm_jsonl import (
    COLUMNS,
    DEFAULT_EXCEL,
    DEFAULT_OUT,
    SKILL_TYPES,
    SPLITS,
    TYPES,
    read_records,
    to_record,
)

MIN_ROWS_PER_TYPE = 30


@pytest.fixture(scope="module")
def records():
    return [json.loads(line) for line in DEFAULT_OUT.read_text(encoding="utf-8").splitlines()]


def test_the_jsonl_is_exactly_what_the_excel_builds_so_nobody_forgot_to_rebuild(records):
    assert records == read_records(DEFAULT_EXCEL)


def test_every_type_has_enough_rows_and_all_three_splits(records):
    by_type = defaultdict(Counter)
    for r in records:
        by_type[r["type"]][r["split"]] += 1

    assert set(by_type) == set(TYPES)
    for intent, splits in by_type.items():
        assert sum(splits.values()) >= MIN_ROWS_PER_TYPE, intent
        assert all(splits[s] > 0 for s in SPLITS), f"{intent} thiếu phần {splits}"


def test_ids_are_unique(records):
    assert len({r["id"] for r in records}) == len(records)


def test_skills_exist_only_for_profile_and_scenario_rows_that_have_a_profile(records):
    with_skill = [r for r in records if r["skill_md"]]

    assert with_skill, "chưa có dòng nào có skill"
    assert {r["type"] for r in with_skill} <= set(SKILL_TYPES)
    assert all(r["input"]["ho_so"] for r in with_skill)


def test_the_same_message_in_the_same_context_never_has_two_different_types(records):
    seen: dict[tuple, set[str]] = defaultdict(set)
    for r in records:
        context = (
            r["input"]["cau_hoi"].lower(),
            r["input"]["cau_he_thong_dang_hoi"],
            json.dumps(r["input"]["ho_so"]),
        )
        seen[context].add(r["type"])

    conflicts = {key: types for key, types in seen.items() if len(types) > 1}
    assert not conflicts, conflicts


def test_scenarios_always_say_which_mode(records):
    for r in records:
        if r["type"] == "kich_ban":
            assert r["output"]["thuc_the"]["che_do"] in ("what_if", "goal_seeking"), r["id"]


# ---- script đổi Excel sang JSONL ----
def row(**overrides):
    base = dict.fromkeys(COLUMNS)
    base |= {"cau_hoi": "tìm trung tâm sửa bài luận ở Sài Gòn", "tap": "train",
             "thuc_the": '{"dich_vu": "SUA_BAI_LUAN", "thanh_pho": "Sài Gòn"}'}  # fmt: skip
    return base | overrides


def check(intent, raw):
    return to_record(intent, raw, f"{intent}, dòng 2", f"{intent}_001")


def test_a_good_row_becomes_a_record():
    record = check("trung_tam", row())

    assert record["output"] == {
        "type": "trung_tam",
        "nhan_phu": None,
        "thuc_the": {"dich_vu": "SUA_BAI_LUAN", "thanh_pho": "Sài Gòn"},
    }
    assert record["input"]["ho_so"] is None and record["skill_md"] is None


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"tap": "dev"}, "tap phải là"),
        ({"cau_hoi": None}, "thiếu cau_hoi"),
        ({"thuc_the": '{"dich_vu": "LAM_VISA_HO"}'}, "mã dịch vụ lạ"),
        ({"thuc_the": '{"thanh_pho": "TP.HCM"}'}, "có trong cau_hoi"),  # phải ghi "Sài Gòn" như học sinh viết
        ({"thuc_the": '{"mau_sac": "đỏ"}'}, "khóa thực thể lạ"),
        ({"thuc_the": "{không phải json"}, "không phải JSON"),
        ({"nhan_phu": "trung_tam"}, "phải xếp sau"),
        (
            {"nhan_phu": "kien_thuc"},
            "phải xếp sau",
        ),  # kiến thức xếp trước trung tâm nên không làm nhãn phụ được
        ({"skill_md": "# skill"}, "không có skill riêng"),
        ({"cau_he_thong_dang_hoi": "Câu hỏi không có trong chat"}, "không khớp câu nào"),
    ],
)
def test_bad_rows_are_rejected_with_the_sheet_and_line(overrides, message):
    with pytest.raises(ValueError, match=message):
        check("trung_tam", row(**overrides))


def test_profile_values_in_entities_must_be_text_from_the_message():
    with pytest.raises(ValueError, match="đoạn chữ"):
        check("trung_tam", row(thuc_the='{"dich_vu": "SUA_BAI_LUAN", "gpa": 8.2}'))


def test_a_profile_answer_without_a_pending_question_must_carry_a_value():
    with pytest.raises(ValueError, match="phải có giá trị hồ sơ"):
        check("mac_dinh", row(cau_hoi="cập nhật SAT giúp em", thuc_the=None))


def test_a_skill_needs_a_profile():
    with pytest.raises(ValueError, match="phải có hồ sơ"):
        check("ho_so", row(skill_md="# skill", thuc_the=None))


def test_a_scenario_needs_a_mode_and_a_change():
    with pytest.raises(ValueError, match="cần thay_doi"):
        check("kich_ban", row(thuc_the='{"che_do": "what_if"}'))


def test_the_form_is_checked_ids_come_from_the_sheet_and_blank_rows_are_skipped(tmp_path):
    wb = Workbook()
    wb.active.title = "trung_tam"
    wb.active.append(list(COLUMNS))
    wb.active.append([row()[c] for c in COLUMNS])
    wb.active.append([None] * len(COLUMNS))
    wb.active.append([row()[c] for c in COLUMNS])
    path = tmp_path / "du_lieu.xlsx"
    wb.save(path)
    assert [r["id"] for r in read_records(path)] == ["trung_tam_001", "trung_tam_002"]

    wb.active.cell(1, len(COLUMNS) + 1, "ghi_chu")
    wb.save(path)
    with pytest.raises(ValueError, match="đúng các cột"):
        read_records(path)

    wb.active.delete_cols(len(COLUMNS) + 1)
    wb.create_sheet("Huong dan")
    wb.save(path)
    with pytest.raises(ValueError, match="mỗi loại một sheet"):
        read_records(path)
